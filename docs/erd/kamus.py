import subprocess, re, os
AKAR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))   # folder akar proyek
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
Q=lambda s: [r.split('|') for r in subprocess.run(['psql','-h','/var/tmp/pg','-U','postgres','-d','spmb_jatimulya04','-AtF|','-c',s],capture_output=True,text=True).stdout.strip().split('\n') if r]
ENT={'pengguna':'Akun login orang tua/wali, panitia SPMB, kepala sekolah, dan admin sekolah',
'periode_spmb':'Jadwal dan daya tampung SPMB per tahun ajaran',
'kuota_jalur':'Persentase kuota setiap jalur pada suatu periode',
'jenis_dokumen':'Master jenis dokumen persyaratan (KK, akta, KTP, dll.)',
'persyaratan_jalur':'Dokumen yang wajib/opsional untuk setiap jalur (tabel asosiatif)',
'calon_murid':'Data pribadi calon murid sesuai Formulir Peserta Didik (Dapodik)',
'orang_tua_wali':'Data ayah, ibu, dan/atau wali calon murid',
'pendaftaran':'Transaksi pendaftaran calon murid pada satu periode dan satu jalur',
'dokumen':'Berkas persyaratan yang diunggah untuk suatu pendaftaran',
'riwayat_status':'Jejak audit setiap perubahan status pendaftaran (diisi trigger)'}
URUT=list(ENT)
D={'id':'Kunci utama (nomor urut otomatis)','pengguna_id':'Akun orang tua/wali pemilik data','nama':'Nama','email':'Alamat email (opsional, unik bila diisi); untuk tautan lupa kata sandi dan dapat dipakai untuk login','username':'Nama pengguna untuk login; boleh nama sendiri dengan spasi, 3–50 karakter, unik tanpa membedakan huruf besar/kecil','no_hp':'Nomor HP/WhatsApp, diawali 08','password_hash':'Kata sandi yang sudah di-hash (bcrypt), bukan teks asli','peran':'Peran pengguna: orang_tua, panitia, kepala_sekolah, admin','dibuat_pada':'Waktu data dibuat','diubah_pada':'Waktu data terakhir diubah',
'tahun_ajaran':'Tahun ajaran, contoh 2026/2027','tgl_mulai_daftar':'Tanggal mulai pendaftaran','tgl_selesai_daftar':'Tanggal akhir pendaftaran','tgl_pengumuman':'Tanggal pengumuman hasil','tgl_daftar_ulang_mulai':'Tanggal mulai daftar ulang','tgl_daftar_ulang_selesai':'Tanggal akhir daftar ulang','tgl_acuan_usia':'Tanggal acuan perhitungan usia (1 Juli tahun berjalan)','jumlah_rombel':'Jumlah rombongan belajar kelas 1','murid_per_rombel':'Jumlah murid per rombel (maks. 28)','daya_tampung':'Kolom terhitung: jumlah_rombel × murid_per_rombel','aktif':'Penanda periode yang sedang berjalan (hanya satu yang aktif)','seleksi_diproses_pada':'Waktu panitia menetapkan hasil seleksi','hasil_diumumkan':'Penanda hasil seleksi sudah dipublikasikan (mengunci pengaturan)','peringkat':'Urutan hasil seleksi dalam jalurnya (usia → jarak → waktu daftar)','daftar_ulang_pada':'Waktu orang tua mengonfirmasi daftar ulang',
'periode_id':'Periode SPMB terkait','jalur':'Jalur pendaftaran: domisili, afirmasi, mutasi','persen':'Persentase kuota jalur dari daya tampung',
'kode':'Kode singkat jenis dokumen, contoh KK, AKTA','jenis_dokumen_id':'Jenis dokumen terkait','wajib':'Apakah dokumen wajib untuk jalur tersebut',
'nik':'Nomor Induk Kependudukan (16 digit)','nisn':'Nomor Induk Siswa Nasional (10 digit, opsional)','nama_lengkap':'Nama lengkap sesuai akta','jenis_kelamin':'L atau P','tempat_lahir':'Tempat lahir','tanggal_lahir':'Tanggal lahir (dasar hitung usia)','no_akta_lahir':'Nomor registrasi akta kelahiran','agama':'Agama','kewarganegaraan':'Kewarganegaraan (default WNI)','alamat':'Alamat jalan/kampung','rt':'RT','rw':'RW','kelurahan':'Kelurahan/desa','kecamatan':'Kecamatan','kode_pos':'Kode pos 5 digit','patokan_rumah':'Patokan/ciri rumah (Dapodik butir 14)','tempat_tinggal':'Tempat tinggal: rumah_sendiri, orang_tua, saudara (butir 15)','moda_transportasi':'Moda transportasi ke sekolah (butir 16)','no_kk':'Nomor Kartu Keluarga (16 digit)','tgl_terbit_kk':'Tanggal terbit KK (syarat domisili ≥ 1 tahun)','jarak_ke_sekolah_km':'Jarak rumah ke sekolah dalam km','anak_ke':'Anak ke-','jumlah_saudara':'Jumlah saudara kandung','no_kks':'Nomor Kartu Keluarga Sejahtera (jalur afirmasi)','penerima_pkh':'Penerima Program Keluarga Harapan','no_kip':'Nomor Kartu Indonesia Pintar','no_kis_pbi':'Nomor KIS Penerima Bantuan Iuran (butir 20)','tinggi_badan_cm':'Tinggi badan dalam cm (butir 21)','berat_badan_kg':'Berat badan dalam kg (butir 22)','waktu_ke_sekolah_menit':'Waktu tempuh ke sekolah dalam menit (butir 24)','berkebutuhan_khusus':'Jenis kebutuhan khusus, bila ada',
'calon_murid_id':'Calon murid terkait','hubungan':'Hubungan: ayah, ibu, wali','tahun_lahir':'Tahun lahir orang tua/wali','pendidikan':'Pendidikan terakhir','pekerjaan':'Pekerjaan','penghasilan':'Rentang penghasilan per bulan',
'nomor_pendaftaran':'Nomor pendaftaran unik, terisi saat diajukan','status':'Status: draf, diajukan, perlu_perbaikan, terverifikasi, diterima, tidak_diterima','usia_bulan':'Usia calon murid (bulan) pada tanggal acuan','perlu_rekomendasi_psikolog':'Usia 5 th 6 bln – < 6 th, wajib rekomendasi psikolog','catatan_panitia':'Catatan panitia untuk perbaikan','diajukan_pada':'Waktu pendaftaran diajukan',
'pendaftaran_id':'Pendaftaran terkait','nama_file_asli':'Nama berkas asli dari pengguna','path_file':'Nama berkas acak di server','tipe_mime':'Tipe berkas (PDF/JPG/PNG)','ukuran_byte':'Ukuran berkas (maks. 2 MB)','status_verifikasi':'Status verifikasi dokumen: menunggu, valid, tidak_valid','catatan':'Catatan','diunggah_pada':'Waktu unggah',
'status_lama':'Status sebelum berubah','status_baru':'Status sesudah berubah','oleh_pengguna_id':'Pengguna yang mengubah status','waktu':'Waktu perubahan status'}
cols=Q("""select c.table_name,c.ordinal_position,c.column_name,
 case when c.data_type='USER-DEFINED' then c.udt_name when c.character_maximum_length is not null then c.data_type||'('||c.character_maximum_length||')' when c.data_type='numeric' and c.numeric_precision is not null then 'numeric('||c.numeric_precision||','||c.numeric_scale||')' else c.data_type end,
 c.is_nullable, coalesce(c.column_default,''), coalesce(c.generation_expression,'')
 from information_schema.columns c join information_schema.tables t using(table_schema,table_name)
 where c.table_schema='spmb' and t.table_type='BASE TABLE' order by 1,2""")
keys=Q("""select tc.table_name,kcu.column_name,tc.constraint_type,coalesce(ccu.table_name,'')
 from information_schema.table_constraints tc join information_schema.key_column_usage kcu using(constraint_schema,constraint_name)
 left join information_schema.constraint_column_usage ccu on tc.constraint_type='FOREIGN KEY' and ccu.constraint_name=tc.constraint_name and ccu.constraint_schema=tc.constraint_schema
 where tc.table_schema='spmb' and tc.constraint_type in ('PRIMARY KEY','FOREIGN KEY','UNIQUE')""")
K={}
for t,c,ty,ref in keys:
    lab={'PRIMARY KEY':'PK','UNIQUE':'UK'}.get(ty,'FK → '+ref)
    K.setdefault((t,c),[]); 
    if lab not in K[(t,c)]: K[(t,c)].append(lab)
chk=Q("""select conrelid::regclass::text, pg_get_constraintdef(oid) from pg_constraint where connamespace='spmb'::regnamespace and contype='c'""")
fks=Q("""select conrelid::regclass::text, confrelid::regclass::text, pg_get_constraintdef(oid), confdeltype from pg_constraint where connamespace='spmb'::regnamespace and contype='f' order by 1""")
enums=Q("select t.typname, string_agg(e.enumlabel, ', ' order by e.enumsortorder) from pg_type t join pg_enum e on e.enumtypid=t.oid group by 1 order by 1")
F='Arial'; H=Font(name=F,bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1D3A6E'); N=Font(name=F,size=10); B=Font(name=F,size=10,bold=True)
th=Side(style='thin',color='999999'); BD=Border(top=th,bottom=th,left=th,right=th)
TF=PatternFill('solid',fgColor='EEF2FA')
def sheet(ws,title,head,rows,widths):
    ws['A1']=title; ws['A1'].font=Font(name=F,size=13,bold=True)
    ws['A2']='Sistem Informasi SPMB Online SDN Jatimulya 04 – Kelompok C (STSI4440). Sumber: skema PostgreSQL schema.sql (skema spmb).'; ws['A2'].font=Font(name=F,size=9,italic=True)
    for j,h in enumerate(head,1):
        c=ws.cell(4,j,h); c.font=H; c.fill=HF; c.border=BD; c.alignment=Alignment(vertical='center',wrap_text=True)
    for i,r in enumerate(rows,5):
        for j,v in enumerate(r,1):
            c=ws.cell(i,j,v); c.font=N; c.border=BD; c.alignment=Alignment(vertical='top',wrap_text=True)
    for j,w in enumerate(widths,1): ws.column_dimensions[chr(64+j)].width=w
    ws.freeze_panes='A5'; ws.auto_filter.ref=f"A4:{chr(64+len(head))}{4+len(rows)}"
wb=Workbook()
cnt={}
for r in cols: cnt[r[0]]=cnt.get(r[0],0)+1
ws=wb.active; ws.title='Daftar Entitas'
rows=[[i,t.upper(),ENT[t],cnt[t]] for i,t in enumerate(URUT,1)]
sheet(ws,'Daftar Entitas',['No','Entitas (tabel)','Keterangan','Jumlah atribut'],rows,[6,24,70,15])
n=5+len(rows); ws.cell(n,2,'Total').font=B; ws.cell(n,4,f'=SUM(D5:D{n-1})').font=B
ws=wb.create_sheet('Kamus Data'); rows=[]
for t in URUT:
    for r in [c for c in cols if c[0]==t]:
        _,pos,c,ty,nl,df,gen=r
        ket=D.get(c,'');
        if c=='nama' and t=='jenis_dokumen': ket='Nama jenis dokumen'
        if c=='nama' and t=='orang_tua_wali': ket='Nama ayah/ibu/wali'
        if c=='nama' and t=='pengguna': ket='Nama lengkap pengguna'
        if c=='aktif' and t=='pengguna': ket='Status akun; false = dinonaktifkan admin dan tidak dapat login'
        if c=='catatan' and t=='riwayat_status': ket='Catatan alasan perubahan status'
        if c=='catatan' and t=='dokumen': ket='Catatan verifikasi dari panitia'
        if c=='id' and t=='calon_murid': pass
        dfl=('GENERATED: '+gen) if gen else re.sub(r"::[a-z_. ]+","",df).replace("nextval('","serial (").replace("_seq'","_seq") 
        ty=ty.replace('character varying','varchar').replace('timestamp with time zone','timestamptz')
        rows.append([t.upper(),int(pos),c,ty,', '.join(K.get((t,c),[])),'Tidak' if nl=='YES' else 'Ya',dfl,ket])
sheet(ws,'Kamus Data (Data Dictionary)',['Entitas','No','Atribut','Tipe data','Kunci','Wajib (NOT NULL)','Default / rumus','Keterangan'],rows,[18,5,26,20,22,10,28,58])
for i,r in enumerate(rows,5):
    if r[1]==1:
        for j in range(1,9): ws.cell(i,j).fill=TF
ws=wb.create_sheet('Relasi')
NAMA={('kuota_jalur','periode_spmb'):'membagi',('pendaftaran','periode_spmb'):'menampung',('calon_murid','pengguna'):'mendaftarkan',('orang_tua_wali','calon_murid'):'memiliki',('pendaftaran','calon_murid'):'diajukan untuk',('dokumen','pendaftaran'):'melampirkan',('dokumen','jenis_dokumen'):'berjenis',('persyaratan_jalur','jenis_dokumen'):'disyaratkan',('riwayat_status','pendaftaran'):'mencatat',('riwayat_status','pengguna'):'mengubah'}
DEL={'a':'NO ACTION','r':'RESTRICT','c':'CASCADE','n':'SET NULL','d':'SET DEFAULT'}
FKI={(a.replace('spmb.',''),b.replace('spmb.','')):(re.search(r'\((\w+)\)',d).group(1),DEL[dt]) for a,b,d,dt in fks}
SYM={'o|':'0..1 (opsional, maks. satu)','|{':'1..N (minimal satu)','o{':'0..N (boleh tidak ada)','||':'1 (tepat satu)'}
rows=[]
for l in open(os.path.join(AKAR, 'docs', 'diagram', 'erd_lengkap.mmd'), encoding='utf-8'):
    m=re.match(r'\s*(\w+) (\|\|)--(\S\S) (\w+) : "(.+)"',l)
    if not m: continue
    P,_,cs,C,nm=m.groups(); fk,od=FKI[(C.lower(),P.lower())]
    rows.append([len(rows)+1,P,'1 (tepat satu)',nm,SYM[cs],C,C+'.'+fk+' → '+P+'.id',od])
sheet(ws,'Relasi Antar-Entitas (sesuai ERD)',['No','Entitas induk','Kardinalitas induk','Nama relasi','Kardinalitas anak','Entitas anak','Kunci tamu (FK)','ON DELETE'],rows,[5,18,14,16,24,18,44,13])
ws=wb.create_sheet('Enum & Aturan')
rows=[['Tipe enumerasi',a,b] for a,b in enums]+[['CHECK',a.replace('spmb.','').upper(),b] for a,b in chk]
rows+=[['Trigger','PENDAFTARAN','fn_catat_riwayat_status: setiap perubahan kolom status otomatis dicatat ke RIWAYAT_STATUS'],
['View','v_rekap_jalur','Rekap per jalur: kuota = ROUND(daya_tampung × persen / 100), jumlah pendaftar, perlu perbaikan, terverifikasi, diterima, tidak diterima, daftar ulang'],
['Indeks unik','PERIODE_SPMB','Hanya satu periode yang boleh aktif = true'],
['Unik gabungan','PENDAFTARAN','UNIQUE(calon_murid_id, periode_id): satu calon murid hanya satu pendaftaran per periode']]
sheet(ws,'Tipe Enumerasi, Batasan (CHECK), Trigger & View',['Jenis','Objek','Definisi'],rows,[16,26,100])
out=os.path.join(AKAR, 'docs', 'erd', 'Kamus Data ERD SPMB SDN Jatimulya 04.xlsx'); wb.save(out); print(len(cols),len(fks),len(chk),len(enums))
